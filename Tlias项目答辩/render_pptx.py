# -*- coding: utf-8 -*-
"""轻量 PPTX 渲染器：读取 python-pptx 生成的形状/文本/图片，用 Pillow 逐页绘制 PNG，用于排版目视校验。"""
import os, sys
from pptx import Presentation
from pptx.util import Emu
from pptx.enum.shapes import MSO_SHAPE_TYPE, MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from PIL import Image, ImageDraw, ImageFont

BASE = os.path.dirname(os.path.abspath(__file__))
PPTX = os.path.join(BASE, "Tlias项目答辩.pptx")
OUTDIR = os.path.join(BASE, "preview")
os.makedirs(OUTDIR, exist_ok=True)
SCALE = 1.0
EMU = 9525.0

FONTS = {
    ("Microsoft YaHei", False): r"C:\Windows\Fonts\msyh.ttc",
    ("Microsoft YaHei", True):  r"C:\Windows\Fonts\msyhbd.ttc",
    ("Consolas", False): r"C:\Windows\Fonts\consola.ttf",
    ("Consolas", True):  r"C:\Windows\Fonts\consolab.ttf",
}
_cache = {}
def font(name, bold, size):
    key = (name, bold, int(size))
    if key not in _cache:
        path = FONTS.get((name, bold)) or FONTS[("Microsoft YaHei", bold)]
        _cache[key] = ImageFont.truetype(path, max(6, int(size)))
    return _cache[key]

def px(v): return int(round(v / EMU * SCALE))

def rgb(color_fmt, default=None):
    try:
        c = color_fmt.rgb
        return (c[0], c[1], c[2])
    except Exception:
        return default

def shape_fill(sh):
    try:
        f = sh.fill
        if f.type == 1:  # solid
            return rgb(f.fore_color)
        if f.type == 3:  # gradient
            try: return rgb(f.gradient_stops[0].color)
            except Exception: return None
    except Exception:
        pass
    return None

def shape_line(sh):
    try:
        if sh.line.fill.type == 1:
            return rgb(sh.line.color), max(1, px(sh.line.width or 12700))
    except Exception:
        pass
    return None, 0

def draw_shape(d, sh):
    x, y = px(sh.left), px(sh.top)
    w, h = px(sh.width), px(sh.height)
    fill = shape_fill(sh)
    lc, lw = shape_line(sh)
    st = None
    try: st = sh.auto_shape_type
    except Exception: pass
    if st == MSO_SHAPE.ISOSCELES_TRIANGLE:
        pts = [(x, y), (x + w, y), (x + w // 2, y + h)]
        if getattr(sh, "rotation", 0) == 180:
            pts = [(x + w, y + h), (x, y + h), (x + w // 2, y)]
        elif getattr(sh, "rotation", 0) == 90:
            pts = [(x, y), (x + w, y + h // 2), (x, y + h)]
        if fill: d.polygon(pts, fill=fill)
    elif st == MSO_SHAPE.OVAL:
        d.ellipse([x, y, x + w, y + h], fill=fill, outline=lc, width=lw or 1)
    elif st == MSO_SHAPE.ROUNDED_RECTANGLE:
        r = min(w, h) // 5
        d.rounded_rectangle([x, y, x + w, y + h], radius=max(2, r), fill=fill,
                            outline=lc, width=lw or 1)
    else:
        d.rectangle([x, y, x + w, y + h], fill=fill, outline=lc, width=lw or 1)

def runs_of(p):
    out = []
    for r in p.runs:
        sz = 15
        try: sz = r.font.size.pt if r.font.size else 15
        except Exception: pass
        bold = bool(r.font.bold)
        name = "Microsoft YaHei"
        try:
            if r.font.name: name = r.font.name
        except Exception: pass
        col = (0x1A, 0x22, 0x33)
        try:
            c = rgb(r.font.color)
            if c: col = c
        except Exception: pass
        out.append((r.text, font(name, bold, sz * SCALE * 96 / 72), col))
    return out

def wrap(runs, maxw):
    """按宽度贪心折行，返回 [[(text,font,color)...], ...]"""
    lines, cur, curw = [], [], 0
    for text, f, col in runs:
        if text == "\n":
            lines.append(cur); cur, curw = [], 0
            continue
        for token in text:
            if token == "\n":
                lines.append(cur); cur, curw = [], 0; continue
            w = f.getlength(token)
            if curw + w > maxw and cur:
                lines.append(cur); cur, curw = [], 0
            cur.append((token, f, col)); curw += w
    lines.append(cur)
    return lines

def draw_text(d, sh):
    tf = sh.text_frame
    x, y = px(sh.left), px(sh.top)
    w, h = px(sh.width), px(sh.height)
    blocks = []
    for p in tf.paragraphs:
        rs = runs_of(p)
        if not rs: continue
        ls = p.line_spacing if isinstance(p.line_spacing, float) else 1.2
        maxsz = max((f.size for _, f, _ in rs), default=15)
        lh = maxsz * SCALE * 96 / 72 * ls
        for line in wrap(rs, w):
            blocks.append((line, lh, p.alignment))
    total = sum(b[1] for b in blocks)
    anchor = tf.vertical_anchor
    if anchor == MSO_ANCHOR.MIDDLE: cy = y + (h - total) / 2
    elif anchor == MSO_ANCHOR.BOTTOM: cy = y + h - total
    else: cy = y
    for line, lh, align in blocks:
        if line:
            lw = sum(f.getlength(t) for t, f, _ in line)
            if align == PP_ALIGN.CENTER: cx = x + (w - lw) / 2
            elif align == PP_ALIGN.RIGHT: cx = x + w - lw
            else: cx = x
            base = cy + (lh - max((f.size for _, f, _ in line), default=15) * SCALE * 96 / 72) * 0.55
            for t, f, col in line:
                d.text((cx, base), t, font=f, fill=col)
                cx += f.getlength(t)
        cy += lh

def walk(shapes):
    for sh in shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            for c in walk(sh.shapes):
                yield c
        else:
            yield sh

def render(path, pages=None):
    prs = Presentation(path)
    Wp, Hp = px(prs.slide_width), px(prs.slide_height)
    files = []
    for i, slide in enumerate(prs.slides, start=1):
        if pages and i not in pages: continue
        img = Image.new("RGB", (Wp, Hp), (255, 255, 255))
        d = ImageDraw.Draw(img)
        for sh in walk(slide.shapes):
            if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                try:
                    im = Image.open(sh.image.blob and __import__("io").BytesIO(sh.image.blob))
                    im = im.convert("RGB").resize((max(1, px(sh.width)), max(1, px(sh.height))))
                    img.paste(im, (px(sh.left), px(sh.top)))
                except Exception as e:
                    d.rectangle([px(sh.left), px(sh.top), px(sh.left) + px(sh.width),
                                 px(sh.top) + px(sh.height)], outline=(200, 60, 60), width=2)
                continue
            if sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
                draw_shape(d, sh)
            if sh.has_text_frame and (sh.text_frame.text or "").strip():
                draw_text(d, sh)
        out = os.path.join(OUTDIR, f"slide_{i:02d}.png")
        img.save(out)
        files.append(out)
    return files

if __name__ == "__main__":
    target = PPTX
    pages = None
    args = sys.argv[1:]
    if args and args[0].lower().endswith(".pptx"):
        target = args[0]; args = args[1:]
    if args:
        pages = set(int(a) for a in args)
    if target != PPTX:
        OUTDIR = os.path.join(BASE, "preview_ref")
        os.makedirs(OUTDIR, exist_ok=True)
    for f in render(target, pages):
        print("rendered", os.path.basename(f))
